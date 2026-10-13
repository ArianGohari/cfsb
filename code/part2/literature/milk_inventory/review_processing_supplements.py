"""Read-only, reproducible extraction of Chen/Shome supplements.

Keep reference aliases, sample species and target-gene species separate.
Never infer detection from a fold change or zero-filled missing p value.
"""

import argparse
from collections import Counter
import csv
import hashlib
import json
from pathlib import Path
import re
import subprocess

import openpyxl

ROOT = Path(__file__).resolve().parents[4]
IMPORT = ROOT / "data/datasets/milk_mirna_inventory/supplement_import_20261009.json"
SH = ROOT / "data/sources/shome_2021_milk_processing_ncrnas/supplements"
CH = ROOT / "data/sources/chen_2010_mirnas_raw_fluid_powdered_milk/supplements"
GROUPS = {
    "Control Group": "raw",
    "Group 1": "heated",
    "Group 2": "homogenized",
    "Group 3": "pasteurized",
    "Group 4": "cream",
    "Group 5": "cultured",
    "Group 6": "sweet",
}


def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument(
        "--output",
        type=Path,
        default=ROOT / "output/part2/milk_inventory/supplement_review",
    )
    args = ap.parse_args()
    out = args.output
    out.mkdir(parents=True, exist_ok=True)
    inputs = {}
    outputs = []
    for r in json.loads(IMPORT.read_text()):
        if "sha256" not in r:
            continue
        p = ROOT / r["path"]
        assert sha(p) == r["sha256"], p
        inputs[r["path"]] = r["sha256"]
    inputs[str(IMPORT.relative_to(ROOT))] = sha(IMPORT)

    def write(name, rows):
        assert rows
        p = out / name
        with p.open("w", newline="") as f:
            w = csv.DictWriter(
                f, fieldnames=list(rows[0]), delimiter="\t", lineterminator="\n"
            )
            w.writeheader()
            w.writerows(rows)
        outputs.append(p)

    pdf = CH / "41422_2010_BFcr201080_MOESM262_ESM.pdf"
    text = subprocess.check_output(["pdftotext", "-layout", str(pdf), "-"]).decode()
    p = out / "chen_s3.txt"
    p.write_text(text)
    outputs.append(p)
    chen = []
    for page, body in enumerate(text.split("\f"), 1):
        for line_no, line in enumerate(body.splitlines(), 1):
            for match in re.finditer(
                r"(?<!\S)(\d+)\s+((?:let-|miR-)\S+)\s+(\d+)\s+(\d+)(?!\S)", line
            ):
                chen.append(
                    {
                        "row_id": int(match[1]),
                        "original_name": match[2],
                        "mature_milk_reads": int(match[3]),
                        "colostrum_reads": int(match[4]),
                        "species": "Bos taurus",
                        "sample_unit": "one pool of 40 cows per lactation group",
                        "DOI": "10.1038/cr.2010.80",
                        "source_path": str(pdf.relative_to(ROOT)),
                        "sha256": sha(pdf),
                        "pdf_page": page,
                        "text_line_on_page": line_no,
                    }
                )
    chen.sort(key=lambda r: r["row_id"])
    assert [r["row_id"] for r in chen] == list(range(1, 246))
    assert len({r["original_name"] for r in chen}) == 245
    cnt = {
        k: sum(r[k] > 0 for r in chen) for k in ["mature_milk_reads", "colostrum_reads"]
    }
    totals = {k: sum(r[k] for r in chen) for k in cnt}
    assert list(cnt.values()) == [213, 230]
    assert list(totals.values()) == [1418136, 1594965]
    write("chen_s3_counts.tsv", chen)

    inventory = []
    for f in sorted(SH.rglob("*.xlsx")):
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        for ws in wb:
            inventory.append(
                {
                    "file": str(f.relative_to(ROOT)),
                    "sheet": ws.title,
                    "rows": ws.max_row,
                    "columns": ws.max_column,
                    "sha256": sha(f),
                }
            )
        wb.close()
    write("shome_workbook_inventory.tsv", inventory)
    f = SH / "12864_2021_7964_MOESM8_ESM/FC_Control (RM) vs. All.xlsx"
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    samples = []
    for i, r in enumerate(wb["Sample Manager"].iter_rows(values_only=True), 1):
        if r[1] not in GROUPS:
            continue
        match = re.search(r"-S([1-4])-", r[0])
        assert match, r
        samples.append(
            {
                "sample": r[0],
                "group": r[1],
                "condition": GROUPS[r[1]],
                "run_label": "S" + match[1],
                "species": "Bos taurus",
                "species_evidence": "Primary paper Methods: dairy-cow bulk milk; not inferred from miRNA aliases",
                "source_path": str(f.relative_to(ROOT)),
                "sheet": "Sample Manager",
                "row": i,
            }
        )
    assert len(samples) == 28 and len({r["sample"] for r in samples}) == 28
    assert Counter(r["group"] for r in samples) == {k: 4 for k in GROUPS}
    write("shome_samples.tsv", samples)
    aliases = []
    for i, r in enumerate(wb["miRNA Table"].iter_rows(values_only=True), 1):
        if not re.fullmatch(r"m\d+", str(r[0])):
            continue
        names = r[1].split("/")
        aliases.append(
            {
                "feature_id": r[0],
                "aliases": r[1],
                "bovine_aliases": ";".join(n for n in names if n.startswith("bta-")),
                "human_aliases": ";".join(n for n in names if n.startswith("hsa-")),
                "source_path": str(f.relative_to(ROOT)),
                "sheet": "miRNA Table",
                "row": i,
            }
        )
    assert len({r["feature_id"] for r in aliases}) == len(aliases) == 3094
    write("shome_feature_aliases.tsv", aliases)
    alias_lookup = {r["feature_id"]: r["aliases"] for r in aliases}
    ws = wb["Fold Change & p-Value"]
    data = list(ws.iter_rows(values_only=True))
    contrasts = []
    for i, r in enumerate(data, 1):
        if not re.fullmatch(r"m\d+", str(r[0])):
            continue
        assert r[1] == alias_lookup[r[0]]
        for c in [2, 5, 8, 11, 14, 17]:
            group = data[1][c]
            assert group in GROUPS
            fc, pvalue, flag = r[c : c + 3]
            assert isinstance(fc, (float, int)) and fc >= 0
            assert (
                pvalue == "NA" or isinstance(pvalue, (float, int)) and 0 <= pvalue <= 1
            )
            contrasts.append(
                {
                    "feature_id": r[0],
                    "group": group,
                    "condition": GROUPS[group],
                    "fold_change": fc,
                    "p_value_as_labeled": pvalue,
                    "comment": flag or "",
                    "interpretation": "Uninterpretable: zero tags in both groups"
                    if flag == "C"
                    else "Low tags; caution"
                    if flag in ["A", "B"]
                    else "No low-tag flag; not a detection count",
                    "source_path": str(f.relative_to(ROOT)),
                    "sheet": ws.title,
                    "fc_cell": f"{openpyxl.utils.get_column_letter(c + 1)}{i}",
                    "p_cell": f"{openpyxl.utils.get_column_letter(c + 2)}{i}",
                    "flag_cell": f"{openpyxl.utils.get_column_letter(c + 3)}{i}",
                }
            )
    assert len(contrasts) == 3094 * 6
    write("shome_reported_contrasts.tsv", contrasts)
    wb.close()
    f = SH / "12864_2021_7964_MOESM5_ESM.xlsx"
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    ws = wb.active
    de = []
    data = list(ws.iter_rows(values_only=True))
    for i, r in enumerate(data, 1):
        if i < 18 or not r[0]:
            continue
        for c, condition in zip(
            [1, 4, 7, 10, 13, 16],
            ["heated", "homogenized", "pasteurized", "cream", "sweet", "cultured"],
        ):
            de.append(
                {
                    "original_name": r[0],
                    "condition": condition,
                    "log2FC": r[c],
                    "pvalue": r[c + 1],
                    "padj": r[c + 2],
                    "source_path": str(f.relative_to(ROOT)),
                    "sheet": ws.title,
                    "cell_log2FC": f"{openpyxl.utils.get_column_letter(c + 1)}{i}",
                    "cell_padj": f"{openpyxl.utils.get_column_letter(c + 3)}{i}",
                }
            )
    write("shome_table5_results.tsv", de)
    wb.close()
    summary = {
        "chen": {
            "table_rows": len(chen),
            "positive_names": cnt,
            "read_totals": totals,
            "count_rule": "positive reads in published S3, no new standardized filtering or sequence harmonization",
        },
        "shome": {
            "samples": len(samples),
            "groups": dict(Counter(r["condition"] for r in samples)),
            "feature_rows_supplement8": len(aliases),
            "features_with_bovine_alias": sum(
                bool(r["bovine_aliases"]) for r in aliases
            ),
            "features_with_human_alias": sum(bool(r["human_aliases"]) for r in aliases),
            "contrasts": len(contrasts),
            "flags": dict(Counter(r["comment"] or "none" for r in contrasts)),
            "table5_rows": len(de) // 6,
            "sample_count_matrix_found": False,
            "limits": "Aliases are not unique mature sequences or proof of detection. Supplement 5 and 8 are distinct result sets; no forced merge. No richness per processing step derivable from these contrast tables.",
        },
    }
    p = out / "summary.json"
    p.write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n")
    outputs.append(p)
    inputs[str(Path(__file__).relative_to(ROOT))] = sha(Path(__file__))
    (out / "manifest.json").write_text(
        json.dumps(
            {
                "inputs": inputs,
                "outputs": {p.name: sha(p) for p in outputs},
                "openpyxl": openpyxl.__version__,
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n"
    )
    print(json.dumps(summary, ensure_ascii=False))


if __name__ == "__main__":
    main()
